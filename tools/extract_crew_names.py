from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


NAME_PATTERN = re.compile(
    r"^(?P<base>[A-Za-z0-9]+(?:_[A-Za-z0-9]+)*)_v(?P<major>[1-9]\d*)(?:_(?P<minor>[1-9]\d*))?$"
)
ORDINARY_NUMBER_PATTERN = re.compile(r"_\d+$")

SPECIAL_SPLIT_GROUPS = {
    "voice_message_driver_engine_stalled": (
        (
            "voice_message_driver_engine_stalled__v1_layers",
            "double",
            (
                "voice_message_driver_engine_stalled_v1_1",
                "voice_message_driver_engine_stalled_v1_2",
                "voice_message_driver_engine_stalled_v1_3",
            ),
        ),
        (
            "voice_message_driver_engine_stalled__v2_v3",
            "single",
            (
                "voice_message_driver_engine_stalled_v2",
                "voice_message_driver_engine_stalled_v3",
            ),
        ),
    ),
}


def natural_name_key(name: str) -> tuple[object, ...]:
    match = NAME_PATTERN.fullmatch(name)
    if match is None:
        return (name.casefold(), name)
    minor = match.group("minor")
    return (
        int(match.group("major")),
        minor is not None,
        int(minor) if minor is not None else 0,
        name,
    )


def classify_exclusion(value: object) -> str:
    if value is None or value == "":
        return "空值"
    if not isinstance(value, str):
        return "非文本单元格"
    if value != value.strip():
        return "包含首尾空白"
    if ORDINARY_NUMBER_PATTERN.search(value):
        return "普通数字结尾"
    if "_v" in value:
        return "歧义或不支持的版本结构"
    return "无可识别版本号或标签文字"


def group_type(names: list[tuple[str, re.Match[str]]]) -> str:
    has_single = any(match.group("minor") is None for _, match in names)
    has_double = any(match.group("minor") is not None for _, match in names)
    if has_single and has_double:
        return "mixed"
    if has_double:
        return "double"
    return "single"


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从 Excel 提取 WT-NameRelay 车组名称组。")
    parser.add_argument("--input", type=Path, required=True, help="输入 .xlsx 文件")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("app/resources/data/crew_name_groups.json"),
        help="生成的 JSON 文件",
    )
    parser.add_argument("--sheet", default="Sheet1", help="工作表名称")
    parser.add_argument("--start-row", type=int, default=3)
    parser.add_argument("--end-row", type=int, default=643)
    parser.add_argument("--expected-names", type=int, default=511)
    parser.add_argument("--expected-groups", type=int, default=151)
    return parser.parse_args()


def main() -> int:
    args = parse_arguments()
    if args.start_row > args.end_row:
        raise ValueError("start-row 不能大于 end-row")
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    try:
        worksheet = workbook[args.sheet]
        groups: dict[str, list[tuple[str, re.Match[str]]]] = defaultdict(list)
        exclusions: list[tuple[int, object, str]] = []
        for row in range(args.start_row, args.end_row + 1):
            value = worksheet.cell(row=row, column=1).value
            match = NAME_PATTERN.fullmatch(value) if isinstance(value, str) else None
            if match is None:
                exclusions.append((row, value, classify_exclusion(value)))
                continue
            groups[match.group("base")].append((value, match))
    finally:
        workbook.close()

    output: dict[str, dict[str, object]] = {}
    all_names: list[str] = []
    for base_name in sorted(groups):
        entries = groups[base_name]
        names = sorted((name for name, _ in entries), key=natural_name_key)
        if len(names) != len(set(names)):
            raise ValueError(f"名称组 {base_name} 存在重复名称")
        output[base_name] = {"type": group_type(entries), "names": names}
        all_names.extend(names)

    for base_name, split_groups in SPECIAL_SPLIT_GROUPS.items():
        original = output.pop(base_name, None)
        if original is None:
            raise ValueError(f"特殊拆分组不存在：{base_name}")
        original_names = set(original["names"])
        split_names = {name for _key, _type, names in split_groups for name in names}
        if original_names != split_names:
            raise ValueError(f"特殊拆分组名称不完整：{base_name}")
        for group_key, group_type_name, names in split_groups:
            output[group_key] = {
                "base_name": base_name,
                "type": group_type_name,
                "names": list(names),
            }

    if len(all_names) != args.expected_names:
        raise ValueError(f"有效名称数量为 {len(all_names)}，期望 {args.expected_names}")
    if len(output) != args.expected_groups:
        raise ValueError(f"名称组数量为 {len(output)}，期望 {args.expected_groups}")
    if len(all_names) != len(set(all_names)):
        raise ValueError("不同名称组之间存在重复名称")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"已生成 {args.output}：{len(all_names)} 个名称，{len(output)} 个名称组。")
    print(f"已排除 {len(exclusions)} 项：")
    for row, value, reason in exclusions:
        print(f"  A{row}: {value!r}（{reason}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
